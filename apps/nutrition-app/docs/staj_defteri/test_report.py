"""Regression checks for the journal sources and its generated Word/PDF layout."""

from copy import deepcopy
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.shared import Cm, Pt

import build_report as report


class JournalTests(unittest.TestCase):
    def test_thirty_days_with_varied_paragraph_counts(self):
        days = report.parse_days()
        self.assertEqual([day[0] for day in days], list(range(1, 31)))
        self.assertEqual({len(day[2]) for day in days}, {3, 4, 5})

    def test_missing_day_is_rejected(self):
        source = (report.BASE / "gunlukler.md").read_text(encoding="utf-8")
        source = source.replace("## 30 |", "## 31 |")
        with patch.object(Path, "read_text", return_value=source):
            with self.assertRaises(ValueError):
                report.parse_days()

    def test_screenshot_metadata_rejects_invalid_records(self):
        valid = {"day": 4, "caption": "Örnek", "capture": "Ekranı çekin.", "layout": "portrait"}
        invalid_lists = [
            {}, [None], [{"day": 4}],
            [valid, valid], [{**valid, "day": 8}, valid],
            [{**valid, "day": 0}],
            [{**valid, "day": 31}], [{**valid, "day": "4"}],
            [{**valid, "day": True}], [{**valid, "day": 4.0}],
            [{**valid, "caption": " "}], [{**valid, "capture": None}],
            [{**valid, "layout": "square"}],
        ]
        for items in invalid_lists:
            with self.subTest(items=items):
                with patch.object(Path, "read_text", return_value=json.dumps(items)):
                    with self.assertRaises(ValueError):
                        report.load_screenshots()

    def test_fifteen_figures_are_sequential(self):
        items = report.load_screenshots()
        self.assertEqual([item["figure"] for item in items.values()], list(range(1, 16)))
        self.assertEqual(sum(item["layout"] == "portrait" for item in items.values()), 14)
        self.assertEqual(items[29]["layout"], "landscape")

    def test_empty_and_confirmed_date_inputs(self):
        info = json.loads((report.BASE / "bilgiler.json").read_text(encoding="utf-8"))
        report.validate_info(info)
        valid = {**info, "gun_tarihleri": [f"{day:02d}/01/2000" for day in range(1, 31)],
                 "baslangic": "01/01/2000", "bitis": "30/01/2000"}
        report.validate_info(valid)
        invalid = deepcopy(valid)
        invalid["gun_tarihleri"][1] = invalid["gun_tarihleri"][0]
        for values in (invalid, {**valid, "gun_tarihleri": valid["gun_tarihleri"][:29]},
                       {**valid, "bitis": "31/01/2000"}):
            with self.subTest(values=values):
                with self.assertRaises(ValueError):
                    report.validate_info(values)

    def test_portrait_layout_has_native_editable_placeholder(self):
        doc = Document()
        item = report.load_screenshots()[4]
        paras = ["First paragraph.", "Second paragraph.", "Third paragraph."]
        report.daily_narrative(doc, paras, item)
        self.assertEqual([p.text for p in doc.paragraphs], paras[:2])
        outer = doc.tables[0]
        self.assertEqual([p.text for p in outer.cell(0, 0).paragraphs], paras[2:])
        box = outer.cell(0, 1).tables[0]
        self.assertEqual(box.cell(0, 0).text, report.screenshot_marker(item))
        self.assertAlmostEqual(box.rows[0].height, Cm(9.6), delta=Pt(1))
        self.assertEqual(outer.cell(0, 1).paragraphs[-1].text, report.screenshot_caption(item))
        self.assertEqual(len(doc.inline_shapes), 0)

    def test_generated_capture_guide_matches_metadata(self):
        items = report.load_screenshots()
        with tempfile.TemporaryDirectory(prefix="staj-guide-test-") as directory:
            with patch.object(report, "BASE", Path(directory)):
                report.write_capture_guide(items)
            text = (Path(directory) / "EKRAN_GORUNTUSU_REHBERI.md").read_text(encoding="utf-8")
        for item in items.values():
            self.assertIn(item["caption"], text)
            self.assertIn(item["capture"], text)
            self.assertIn(f"{item['day']:02d} / {item['day'] + 2}", text)

    def test_generated_word_preserves_twelve_point_narrative(self):
        def paragraphs(container):
            yield from container.paragraphs
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from paragraphs(cell)

        doc = Document(report.BASE / f"{report.STEM}.docx")
        by_text = {p.text: p for p in paragraphs(doc)}
        for number, _, paras in report.parse_days():
            for text in paras:
                with self.subTest(day=number, paragraph=text[:30]):
                    self.assertIn(text, by_text)
                    self.assertTrue(all(run.font.size == Pt(12) for run in by_text[text].runs if run.text))

    def test_generated_pdf_matches_validation_record(self):
        actual = report.validate_pdf(report.BASE / f"{report.STEM}.pdf",
                                     report.parse_days(), report.load_screenshots())
        recorded = json.loads((report.BASE / "validation.json").read_text(encoding="utf-8"))
        self.assertEqual(actual, recorded)

    def test_package_documentation_links_resolve(self):
        for source in report.BASE.glob("*.md"):
            text = source.read_text(encoding="utf-8")
            for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", text):
                if "://" in target or target.startswith("#"):
                    continue
                with self.subTest(source=source.name, target=target):
                    self.assertTrue((source.parent / target.split("#", 1)[0]).exists())


if __name__ == "__main__":
    unittest.main()
