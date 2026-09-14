"""Build the original app-only Turkish internship draft; never copy sample assets.

Requires python-docx, LibreOffice, pdfinfo and pdftotext. Outputs are generated
from gunlukler.md + bilgiler.json. Source changes must be made there, not in PDF.
Personal fields remain placeholders unless the operator explicitly fills them.
"""

from __future__ import annotations

import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from statistics import mean
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
STEM = "OpenNutri_Staj_Defteri_30_Gun"
FONT = "Arial"


def xml(parent, tag, **attrs):
    element = OxmlElement(f"w:{tag}")
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), str(value))
    parent.append(element)
    return element


def font(run, size=12, bold=False, color="000000"):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    xml(props, "lang", val="tr-TR")
    return run


def paragraph(container, text="", *, size=12, bold=False,
              align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=11):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.widow_control = True
    font(p.add_run(text), size, bold)
    return p


def cell_text(cell, text, *, size=10, bold=False, center=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    font(p.add_run(text), size, bold)


def table_borders(table, color="222222"):
    props = table._tbl.tblPr
    borders = xml(props, "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        xml(borders, edge, val="single", sz="6", color=color)
    margins = xml(props, "tblCellMar")
    for edge in ("top", "bottom"):
        xml(margins, edge, w="70", type="dxa")
    for edge in ("left", "right"):
        xml(margins, edge, w="110", type="dxa")
    for row in table.rows:
        xml(row._tr.get_or_add_trPr(), "cantSplit")


def section_layout(section, *, cover=False):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.top_margin = Cm(3.6 if not cover else 2.8)
    section.bottom_margin = Cm(3.1)
    section.header_distance = Cm(1.65)
    section.footer_distance = Cm(1.45)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for old in list(section._sectPr.findall(qn("w:pgBorders"))):
        section._sectPr.remove(old)
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        xml(borders, edge, val="single", sz="6", space="25", color="222222")
    section._sectPr.append(borders)


def empty_header_footer(section):
    for container in (section.header, section.footer):
        p = container.paragraphs[0]
        p.text = ""
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        font(p.add_run(""), 1)


def daily_frame(section, info, number):
    empty_header_footer(section)
    header = section.header.add_table(rows=1, cols=2, width=Cm(17.3))
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    header.columns[0].width = Cm(12.4)
    header.columns[1].width = Cm(4.9)
    cell_text(header.cell(0, 0), f"Çalışılan Kurum: {info['kurum']}", size=10.5, bold=True)
    cell_text(header.cell(0, 1), f"Sayfa No: {number:02d}", size=11, bold=True, center=True)
    table_borders(header)
    footer = section.footer.add_table(rows=2, cols=3, width=Cm(17.3))
    footer.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer.autofit = False
    for col, width in zip(footer.columns, (5.3, 6.8, 5.2)):
        col.width = Cm(width)
    date = info['gun_tarihleri'][number - 1] if info['gun_tarihleri'] else "[GG/AA/YYYY]"
    for i, label in enumerate(("Çalışılan Tarih", "Onaylayanın Adı Soyadı", "İmza ve Mühür")):
        cell_text(footer.cell(0, i), label, size=9, bold=True, center=True)
    for i, value in enumerate((date, info["sorumlu"], "")):
        cell_text(footer.cell(1, i), value, size=10, center=True)
    footer.rows[1].height = Cm(0.7)
    table_borders(footer)


def cover_page(doc, info):
    section_layout(doc.sections[0], cover=True)
    empty_header_footer(doc.sections[0])
    for text, size in (("T.C.", 13), (info['universite'], 16), (info['fakulte'], 12)):
        paragraph(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=15)
    paragraph(doc, "", after=24)
    paragraph(doc, "STAJ DEFTERİ", size=25, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    paragraph(doc, "OpenNutri", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=9)
    paragraph(doc, "Sesli Besin Kaydı ve Beslenme Takibi\nMobil Uygulaması", size=14,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=32)
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(5)
    table.columns[1].width = Cm(12.3)
    for label, value in (("Bölümü", info['bolum']), ("Öğrenci Numarası", info['ogrenci_no']),
                         ("Adı Soyadı", info['ogrenci_adi']), ("Çalışılan Kurum", info['kurum']),
                         ("Çalışma Devresi", "Yaz Stajı"),
                         ("Tarih Aralığı", f"{info['baslangic']} – {info['bitis']}")):
        cells = table.add_row().cells
        cell_text(cells[0], label, size=11, bold=True)
        cell_text(cells[1], value, size=11)
    table_borders(table)
    paragraph(doc, "", after=15)
    paragraph(doc, "30 günlük içerik taslağı • Tarih ve katkı doğrulaması bekliyor",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    paragraph(doc, "Kurumun resmî kapak ve onay formu varsa bu kapağın yerine kullanılmalıdır.",
              size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)


def information_page(doc, info):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section_layout(section, cover=True)
    empty_header_footer(section)
    paragraph(doc, "PRATİK ÇALIŞMA BİLGİLERİ", size=15, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(6)
    table.columns[1].width = Cm(11.3)
    for label, value in (
        ("Öğrenci Adı Soyadı", info['ogrenci_adi']), ("Öğrenci Numarası / Bölümü", f"{info['ogrenci_no']} / {info['bolum']}"),
        ("Kurum Adı", info['kurum']), ("Kurum Adresi", info['kurum_adresi']),
        ("Sorumlu Adı Soyadı", info['sorumlu']), ("Sorumlu Unvanı", info['sorumlu_unvani']),
        ("Başlangıç / Bitiş Tarihi", f"{info['baslangic']} / {info['bitis']}"),
        ("Günlük İçerik Sayısı", "30 — fiilî çalışma günleriyle eşleştirilecek"),
    ):
        cells = table.add_row().cells
        cell_text(cells[0], label, size=10.5, bold=True)
        cell_text(cells[1], value, size=10.5)
    table_borders(table)
    paragraph(doc, "", after=9)
    paragraph(doc, "Belgenin kapsamı", size=12, bold=True, after=7)
    paragraph(doc, "Bu defter yalnızca OpenNutri mobil uygulamasını anlatır. Mevcut besin ve yapay zekâ servisleri, uygulamanın kullandığı dış bileşenler olarak ele alınmıştır. Veri kaynağının oluşturulması ve diğer proje çalışmalarına ilişkin geliştirme iddiası içermez.", size=10.5, after=12)
    paragraph(doc, "Teslim öncesi doğrulama", size=12, bold=True, after=7)
    paragraph(doc, "Birinci tekil şahısla yazılan günlükler, öğrenci tarafından doğrulanacak anlatım önerileridir. Yapılan işi açıklayan kod ve testler, o işi öğrencinin yaptığını veya belirtilen sırada çalıştığını tek başına kanıtlamaz. Öğrenci, geliştirme, entegrasyon, test ve inceleme ifadelerini gerçek katkısına göre düzeltmeli; tarihleri doğrulamalıdır. Sayısal örnekler kodu açıklayan kontrollü verilerdir. Metin yapay zekâ desteğiyle hazırlanmış; kurum bilgileri ve imzalar boş bırakılmıştır.", size=10.5, after=12)
    paragraph(doc, "Onay alanı — ilgili yetkili tarafından doldurulur", size=10.5, bold=True, after=12)
    paragraph(doc, "Adı Soyadı / Unvanı: ........................................................................\n\nTarih: ............................       İmza ve Mühür: ............................", size=10, align=WD_ALIGN_PARAGRAPH.LEFT)


def technical_table(doc, title, rows):
    p = paragraph(doc, title, size=9.5, bold=True, after=5)
    p.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(5)
    table.columns[1].width = Cm(12.3)
    for label, value in rows:
        cells = table.add_row().cells
        cell_text(cells[0], label, size=9, bold=True)
        cell_text(cells[1], value, size=9)
    table_borders(table, color="777777")


def parse_days():
    text = (BASE / "gunlukler.md").read_text(encoding="utf-8")
    chunks = re.split(r"^## (\d+) \| (.+)$", text, flags=re.MULTILINE)
    days = []
    for i in range(1, len(chunks), 3):
        number, title, body = chunks[i:i + 3]
        days.append((int(number), title, body.strip().split("\n\n")))
    if [n for n, _, _ in days] != list(range(1, 31)):
        raise ValueError("Exactly 30 numbered daily entries are required")
    if any(len(paras) != 4 for _, _, paras in days):
        raise ValueError("Each detailed daily entry must have four narrative paragraphs")
    return days


def validate_info(info):
    dates = info["gun_tarihleri"]
    if not isinstance(dates, list) or len(dates) not in (0, 30):
        raise ValueError("gun_tarihleri must be empty or contain 30 confirmed dates")
    if dates:
        parsed = [datetime.strptime(d, "%d/%m/%Y").date() for d in dates]
        if any(a >= b for a, b in zip(parsed, parsed[1:])):
            raise ValueError("Confirmed daily dates must be unique and chronological")
        for field, value in (("baslangic", dates[0]), ("bitis", dates[-1])):
            if info[field] != value:
                raise ValueError(f"{field} must match the confirmed daily date range")


def validate_pdf(pdf_path, days):
    metadata = subprocess.check_output(["pdfinfo", str(pdf_path)], text=True)
    count = int(re.search(r"^Pages:\s+(\d+)", metadata, re.MULTILINE)[1])
    raw = subprocess.check_output(["pdftotext", "-layout", str(pdf_path), "-"], text=True)
    pages = [p for p in raw.split("\f") if p.strip()]
    if count != 32 or len(pages) != 32:
        raise ValueError(f"Expected cover + information + 30 pages, got {count}/{len(pages)}")
    daily_counts = []
    for (number, title, paras), page in zip(days, pages[2:]):
        normalized = " ".join(page.split())
        required = [f"Sayfa No: {number:02d}", title, "Çalışılan Tarih", "İmza ve Mühür"]
        if not all(text in normalized for text in required):
            raise ValueError(f"Day {number} missing title or frame")
        for para in paras:
            if " ".join(para.split()) not in normalized:
                raise ValueError(f"Day {number} text missing, split across pages, or corrupted")
        if "\ufffd" in page:
            raise ValueError(f"Day {number} has replacement characters")
        daily_counts.append(len(normalized.split()))
    bbox = subprocess.check_output(["pdftotext", "-bbox", str(pdf_path), "-"], text=True)
    tree = ET.fromstring(bbox)
    ns = {"x": "http://www.w3.org/1999/xhtml"}
    for index, page in enumerate(tree.findall(".//x:page", ns), start=1):
        width, height = float(page.attrib["width"]), float(page.attrib["height"])
        for word in page.findall(".//x:word", ns):
            coords = [float(word.attrib[k]) for k in ("xMin", "yMin", "xMax", "yMax")]
            if coords[0] < 25 or coords[1] < 25 or coords[2] > width - 25 or coords[3] > height - 25:
                raise ValueError(f"Page {index} text outside printable frame")
    body_counts = [sum(len(p.split()) for p in ps) for _, _, ps in days]
    return {"pdf_pages": count, "daily_pages": 30,
            "body_words": sum(body_counts),
            "daily_body_words_min": min(body_counts),
            "daily_body_words_max": max(body_counts),
            "paragraphs_per_day": 4,
            "daily_words_with_frames_and_tables": sum(daily_counts),
            "daily_mean_words": round(mean(daily_counts), 1),
            "one_day_per_page": True, "all_source_paragraphs_present": True,
            "all_text_inside_printable_frame": True}


def main():
    for exe in ("libreoffice", "pdfinfo", "pdftotext"):
        if not shutil.which(exe):
            raise SystemExit(f"Required dependency missing: {exe}")
    info = json.loads((BASE / "bilgiler.json").read_text(encoding="utf-8"))
    validate_info(info)
    days = parse_days()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.22
    doc.core_properties.author = "OpenNutri"
    doc.core_properties.title = "OpenNutri Mobil Uygulaması — Staj Defteri Taslağı"
    doc.core_properties.subject = "30 günlük uygulama odaklı teknik içerik; tarih ve katkı doğrulanmalıdır"
    doc.core_properties.keywords = "Flutter, Dart, Android, mobil uygulama, staj taslağı"
    cover_page(doc, info)
    information_page(doc, info)
    for number, title, paras in days:
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        section_layout(section)
        daily_frame(section, info, number)
        title_p = paragraph(doc, title, size=12.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=16)
        title_p.paragraph_format.keep_with_next = True
        for text in paras:
            paragraph(doc, text, size=12, after=13)
        if number == 3:
            technical_table(doc, "Tablo 1. Mobil uygulamadaki sorumluluk ayrımı", [
                ("Modeller", "Besin, günlük kaydı, hedef ve profil bilgileri"),
                ("Denetleyici / ekranlar", "Ortak durum ve kullanıcı etkileşimleri"),
                ("Servisler", "Yerel saklama ve mevcut HTTP servislerine erişim"),
            ])
        if number == 22:
            technical_table(doc, "Tablo 2. Araç takımından kayıt ekranına geçiş", [
                ("1. Araç takımı", "Mikrofon düğmesi, sesli kayıt Intent işlemini başlatır."),
                ("2. MainActivity", "İstek yakalanır ve bir kez tüketilecek biçimde tutulur."),
                ("3. MethodChannel", "Flutter mevcut sesli kayıt ekranını açar."),
            ])
    docx_path = BASE / f"{STEM}.docx"
    doc.save(docx_path)
    with tempfile.TemporaryDirectory(prefix="opennutri-staj-export-") as profile:
        subprocess.run(["libreoffice", f"-env:UserInstallation={Path(profile).as_uri()}",
                        "--headless", "--convert-to", "pdf", "--outdir", str(BASE),
                        str(docx_path)], check=True, timeout=120)
    report = validate_pdf(BASE / f"{STEM}.pdf", days)
    (BASE / "validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
